"""Extractor for DOCX files using python-docx."""

from __future__ import annotations

import logging
from pathlib import Path

from llmfilesearch.extractors.base import BaseExtractor, ExtractionResult

logger = logging.getLogger(__name__)


class DocxExtractor(BaseExtractor):
    """Extract text and metadata from DOCX files."""

    def extract(self, file_path: Path) -> ExtractionResult:
        try:
            from docx import Document

            doc = Document(str(file_path))

            metadata = {}
            core_props = doc.core_properties
            if core_props.title:
                metadata["title"] = core_props.title
            if core_props.author:
                metadata["author"] = core_props.author
            if core_props.subject:
                metadata["subject"] = core_props.subject
            if core_props.created:
                metadata["created"] = core_props.created.isoformat()
            if core_props.modified:
                metadata["modified"] = core_props.modified.isoformat()

            # Extract paragraphs
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            text = "\n\n".join(paragraphs)

            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        text += "\n" + row_text

            text = self._truncate(text)

            if not text.strip():
                return ExtractionResult(
                    metadata=metadata, status="metadata_only"
                )

            return ExtractionResult(text=text, metadata=metadata)

        except Exception as e:
            logger.warning("Failed to extract DOCX %s: %s", file_path, e)
            return ExtractionResult(status="failed", error=str(e))
