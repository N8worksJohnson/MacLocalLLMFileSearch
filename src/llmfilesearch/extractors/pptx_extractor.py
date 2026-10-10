"""Extractor for PPTX files using python-pptx."""

from __future__ import annotations

import logging
from pathlib import Path

from llmfilesearch.extractors.base import BaseExtractor, ExtractionResult

logger = logging.getLogger(__name__)


class PptxExtractor(BaseExtractor):
    """Extract text and metadata from PowerPoint files."""

    def extract(self, file_path: Path) -> ExtractionResult:
        try:
            from pptx import Presentation

            prs = Presentation(str(file_path))

            metadata = {
                "slide_count": len(prs.slides),
            }

            core_props = prs.core_properties
            if core_props.title:
                metadata["title"] = core_props.title
            if core_props.author:
                metadata["author"] = core_props.author
            if core_props.subject:
                metadata["subject"] = core_props.subject
            if core_props.created:
                metadata["created"] = core_props.created.isoformat()

            # Extract text from all slides
            text_parts = []
            for i, slide in enumerate(prs.slides, 1):
                slide_texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            para_text = paragraph.text.strip()
                            if para_text:
                                slide_texts.append(para_text)
                    if shape.has_table:
                        for row in shape.table.rows:
                            row_text = " | ".join(
                                cell.text.strip()
                                for cell in row.cells
                                if cell.text.strip()
                            )
                            if row_text:
                                slide_texts.append(row_text)

                if slide_texts:
                    text_parts.append(f"Slide {i}:\n" + "\n".join(slide_texts))

            text = "\n\n".join(text_parts)
            text = self._truncate(text)

            if not text.strip():
                return ExtractionResult(
                    metadata=metadata, status="metadata_only"
                )

            return ExtractionResult(text=text, metadata=metadata)

        except Exception as e:
            logger.warning("Failed to extract PPTX %s: %s", file_path, e)
            return ExtractionResult(status="failed", error=str(e))
